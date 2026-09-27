import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml import parse_xml
import lxml.etree

from core.tokens import get_preset_theme
from engine.components import M3Components

theme = get_preset_theme("electric_violet", is_dark=False)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Test 4 cases side by side:
# Case 1: current code (nudge = Pt(3.0), default shadow from pptx)
# Case 2: nudge = Pt(4.5), default shadow
# Case 3: nudge = Pt(4.5), shadow REMOVED (effectRef idx=0, effectLst empty)
# Case 4: nudge = Pt(5.0), shadow REMOVED (effectRef idx=0, effectLst empty)

cases = [
    ("1: Nudge 3.0pt (Current)", Pt(3.0), False),
    ("2: Nudge 4.5pt (Shadow)", Pt(4.5), False),
    ("3: Nudge 4.5pt (NoShadow)", Pt(4.5), True),
    ("4: Nudge 5.0pt (NoShadow)", Pt(5.0), True),
]

items = [
    "单 Prompt 职责过载导致严重幻觉",
    "长链路执行缺乏纠错与自反思闭环",
    "跨领域多工具调用冲突与调度混乱"
]

for col_idx, (label, nudge, remove_shadow) in enumerate(cases):
    left_x = 0.8 + col_idx * 3.0
    
    # Header
    tb_hdr = slide.shapes.add_textbox(Inches(left_x), Inches(0.4), Inches(2.8), Inches(0.4))
    tb_hdr.text_frame.text = label
    tb_hdr.text_frame.paragraphs[0].font.size = Pt(11)
    tb_hdr.text_frame.paragraphs[0].font.bold = True
    
    # Card
    card = M3Components.create_card(
        slide, theme,
        left=Inches(left_x), top=Inches(0.9),
        width=Inches(2.8), height=Inches(2.5),
        variant="filled", color_role="surface_container_high", corner_dp=16, padding=0.15
    )
    
    cur_y = card.inner_bounds.top
    tb_h_emu = Inches(0.32)
    b_size_emu = Inches(0.18)
    spacing_emu = Inches(0.38)
    badge_offset_y = (tb_h_emu - b_size_emu) // 2 - nudge

    for text in items:
        badge = M3Components.create_icon_container(
            slide, theme,
            left=card.inner_bounds.left,
            top=cur_y + badge_offset_y,
            size=0.18,
            icon_name="close",
            color_role="surface_container_highest",
            icon_color_role="outline",
            shape_type="circle"
        )
        
        if remove_shadow:
            # remove shadow by setting effectRef to 0 and adding empty effectLst
            spPr = badge._sp.spPr
            for child in list(spPr):
                if child.tag.endswith('effectLst'):
                    spPr.remove(child)
            spPr.append(parse_xml('<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
            style = badge._sp.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
            if style is not None:
                effectRef = style.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
                if effectRef is not None:
                    effectRef.set('idx', '0')
        
        tb = slide.shapes.add_textbox(
            card.inner_bounds.left + Inches(0.28),
            cur_y,
            card.inner_bounds.width - Inches(0.28),
            tb_h_emu
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(10.5)
        p.font.name = "Noto Sans SC"
        p.font.color.rgb = theme.rgb("on_surface_variant")
        
        cur_y += spacing_emu

os.makedirs("tests/output", exist_ok=True)
prs.save("tests/output/test_shadow_align.pptx")
print("Saved test_shadow_align.pptx")
