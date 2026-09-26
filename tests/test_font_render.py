from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Test 1: Title with Segoe UI Variable Display for both latin and ea
tb1 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(1.5))
p1 = tb1.text_frame.paragraphs[0]
r1 = p1.add_run()
r1.text = "测试大标题：智能云原生架构演进 2026 (Segoe UI Variable Display)"
r1.font.size = Pt(28)
r1.font.bold = True
rPr1 = r1._r.get_or_add_rPr()
rPr1.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))
rPr1.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))

# Test 2: Subtitle / Small title with Segoe UI Variable Display
tb2 = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.0))
p2 = tb2.text_frame.paragraphs[0]
r2 = p2.add_run()
r2.text = "小标题测试：28dp Bento 容器与双强调色感知架构 (Segoe UI Variable Display 18pt)"
r2.font.size = Pt(18)
r2.font.bold = True
rPr2 = r2._r.get_or_add_rPr()
rPr2.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))
rPr2.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))

# Test 3: Body text with Noto Sans SC (14pt)
tb3 = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11), Inches(1.0))
p3 = tb3.text_frame.paragraphs[0]
r3 = p3.add_run()
r3.text = "正文描述测试：基于 Google HCT 感知色彩与多模态端侧大模型，无需联网即可毫秒级理解用户复杂意图与上下文。(Noto Sans SC 14pt)"
r3.font.size = Pt(14)
rPr3 = r3._r.get_or_add_rPr()
rPr3.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Text"/>'))
rPr3.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Noto Sans SC"/>'))

# Test 4: Body text with Noto Sans SC (12pt bullets)
tb4 = slide.shapes.add_textbox(Inches(1), Inches(5.3), Inches(11), Inches(1.5))
p4 = tb4.text_frame.paragraphs[0]
r4 = p4.add_run()
r4.text = "• 清单子项测试：DrawingML 纯净编译，跨设备字体渲染永不走样，支持 PowerPoint 自由拖拽编辑 (Noto Sans SC 12pt)"
r4.font.size = Pt(12)
rPr4 = r4._r.get_or_add_rPr()
rPr4.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Text"/>'))
rPr4.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Noto Sans SC"/>'))

os.makedirs("tests/output", exist_ok=True)
prs.save("tests/output/test_font_render.pptx")
print("Saved tests/output/test_font_render.pptx")
