"""
Full Deck Integration Test for Phase 3: 8 Core Slide Archetypes.
Generates an 8-slide presentation covering all archetypes with consistent theme and styling.
"""

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches

from core.tokens import get_preset_theme
from archetypes import (
    render_cover_hero,
    render_bento_grid,
    render_process_roadmap,
    render_dual_contrast,
    render_kpi_metrics,
    render_feature_list,
    render_quote_takeaway,
    render_team_showcase,
)

def safe_save_pptx(prs, target_path):
    os.makedirs(os.path.dirname(os.path.abspath(target_path)), exist_ok=True)
    try:
        prs.save(target_path)
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        for i in range(1, 20):
            cand = f"{base}_v{i}{ext}"
            try:
                prs.save(cand)
                print(f"Notice: {target_path} is locked by PowerPoint. Saved to {cand} instead.")
                return cand
            except PermissionError:
                continue
        raise PermissionError(f"Could not save {target_path} in any version.")

def build_showcase_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    theme = get_preset_theme("indigo_coral", is_dark=False)

    print("Rendering 8 Slide Archetypes...")

    # 1. Slide 1: Cover Hero
    s1 = prs.slides.add_slide(blank_layout)
    render_cover_hero(s1, theme, {
        "category": "PRODUCT ARCHITECTURE 2026",
        "title": "下一代智能工作流：\nMaterial 3 Expressive 架构演进",
        "subtitle": "融合 Google HCT 感知色彩、Bento 便当盒分舱与 100% 原生矢量 DrawingML 编译管线",
        "speaker": "AI & Systems Architecture Team",
        "date": "2026 年秋季发布",
        "org": "Google DeepMind Advanced Engineering",
        "hero_highlight": "M3 EXP-2026",
        "hero_badge_text": "全场景演示系统工业化重构"
    })
    print("  [OK] Archetype 01: Cover Hero rendered")

    # 2. Slide 2: Bento Grid Dashboard
    s2 = prs.slides.add_slide(blank_layout)
    render_bento_grid(s2, theme, {
        "category": "SYSTEM OVERVIEW",
        "title": "全景看板：Bento 便当盒模块化架构",
        "subtitle": "依循 28dp 核心大圆角与呼吸感间距，构建秩序井然的多舱信息呈现空间",
        "hero_card": {
            "chip": "CORE PILLAR 01", "icon": "auto_awesome",
            "title": "动态感知色彩与双强调色", "subtitle": "Dual-Chroma Dynamic Palette",
            "desc": "基于 Google MCU 算法推导和谐高对比度的 Primary 与 Tertiary 碰撞，使核心数据跃然而出。",
            "metric_val": "86.4%", "metric_label": "视觉注意聚焦点转化率", "trend": "+3.4x 跃迁"
        },
        "card_top": {
            "icon": "grid_view", "title": "BentoGrid 空间算子",
            "desc": "卡片内外边距严格依循 16/24/28dp 黄金律，告别传统死板直角表格。",
            "chips": ["28dp 容器", "9999px 胶囊", "色调浸润底色"]
        },
        "card_bot_left": {
            "icon": "bolt", "stat": "+4.8x",
            "title": "表现力跃迁", "desc": "有机花瓣徽章打破平直呆板。"
        },
        "card_bot_right": {
            "icon": "verified", "title": "原生矢量交付",
            "bullets": ["DrawingML 纯净编译", "抗字体降级体系", "100% 自由编辑修改"]
        },
        "footer_right": "Page 02 / 08"
    })
    print("  [OK] Archetype 02: Bento Grid Dashboard rendered")

    # 3. Slide 3: Process Roadmap / Timeline
    s3 = prs.slides.add_slide(blank_layout)
    render_process_roadmap(s3, theme, {
        "category": "PROJECT ROADMAP",
        "title": "系统演进全景路线图与里程碑",
        "subtitle": "四阶段协同推进：从算法底座整合、表现力组件研发，到完整业务闭环落地",
        "phases": [
            {
                "phase_num": "01", "name": "官方资产集成", "time": "Q1 · 已完成", "status": "COMPLETED",
                "icon": "check", "deliverables": ["Google MCU 算法接入", "33 款矢量图标缓存", "HCT 动态色阶推导"],
                "kpi_chip": "基建 100% 就绪"
            },
            {
                "phase_num": "02", "name": "表现力组件引擎", "time": "Q2 · 已完成", "status": "COMPLETED",
                "icon": "bolt", "deliverables": ["28dp 核心大圆角规范", "9999px 全胶囊药丸算子", "中西文抗降级排版栈"],
                "kpi_chip": "组件引擎就绪"
            },
            {
                "phase_num": "03", "name": "8大核心版式库", "time": "Q3 · 进行中", "status": "ACTIVE",
                "icon": "layers", "deliverables": ["覆盖全场景版式矩阵", "自适应数据注入管道", "氛围光晕与统一底色"],
                "kpi_chip": "8 套版式全覆盖"
            },
            {
                "phase_num": "04", "name": "Skill 封装与交付", "time": "Q4 · 目标", "status": "UPCOMING",
                "icon": "rocket", "deliverables": ["自然语言大纲解析器", "全局 SKILL.md 注册", "端到端一句话生成"],
                "kpi_chip": "工业化落地"
            }
        ],
        "footer_right": "Page 03 / 08"
    })
    print("  [OK] Archetype 03: Process Roadmap rendered")

    # 4. Slide 4: Dual-Contrast Comparison
    s4 = prs.slides.add_slide(blank_layout)
    render_dual_contrast(s4, theme, {
        "category": "ARCHITECTURE COMPARISON",
        "title": "传统演示方案 vs M3 Expressive 深度比对",
        "subtitle": "对比传统沉闷设计局限与下一代表现力架构的体验跃迁",
        "left_plan": {
            "chip": "TRADITIONAL BASELINE",
            "title": "传统低对比度设计局限",
            "desc": "单一冷灰缺乏视觉聚焦，受众阅读易产生认知疲劳。",
            "points": [
                "单调沉闷的四角直角与表格排版",
                "黑重脏阴影破坏画面透气呼吸感",
                "中文字体在跨端时易降级退化为宋体"
            ],
            "metric_val": "34.5%",
            "metric_label": "信息瞬时记忆留存率",
            "trend": "传统基准参考"
        },
        "right_plan": {
            "chip": "M3 EXPRESSIVE EVOLVED",
            "title": "高活力双强调色感知架构",
            "desc": "基于 HCT 算法引入双强调色，重塑高对比度视觉层次。",
            "points": [
                "28dp 核心大圆角与 9999px 全胶囊药丸",
                "Surface Container 纯净色阶代替脏黑阴影",
                "DrawingML 双字族注入，现代几何无衬线抗降级"
            ],
            "metric_val": "86.4%",
            "metric_label": "核心决策信息传达与留存率",
            "trend": "+2.5x 跃迁"
        },
        "footer_right": "Page 04 / 08"
    })
    print("  [OK] Archetype 04: Dual Contrast rendered")

    # 5. Slide 5: KPI Metrics Dashboard
    s5 = prs.slides.add_slide(blank_layout)
    render_kpi_metrics(s5, theme, {
        "category": "PERFORMANCE BENCHMARKS",
        "title": "核心业务运行指标与成果大屏",
        "subtitle": "量化衡量系统在高可用、推理时延、交付吞吐与用户满意度上的全面突破",
        "metrics": [
            {
                "value": "99.98%", "label": "系统高可用性 (SLA)", "trend": "+0.15% 稳定达标",
                "icon": "shield", "desc": "多活容灾与自动化运维保障", "highlight": False
            },
            {
                "value": "12.5 ms", "label": "端到端推理时延", "trend": "-64% 延迟压缩",
                "icon": "speed", "desc": "底层计算图融合与量化加速", "highlight": True
            },
            {
                "value": "4.8x", "label": "全链路研发吞吐", "trend": "+380% 交付提速",
                "icon": "rocket", "desc": "组件化管线带来效能跃迁", "highlight": False
            },
            {
                "value": "86.4%", "label": "用户体验满意度 (CSAT)", "trend": "+24.6% 提升",
                "icon": "star", "desc": "全新 M3 表现力设计语言获高度好评", "highlight": False
            }
        ],
        "takeaway_badge": "KEY TAKEAWAY",
        "takeaway_text": "量化成果洞察：通过全链路架构现代化改造，关键 SLA 指标均超预期达成，为企业级规模化落地提供了高韧性支撑。",
        "footer_right": "Page 05 / 08"
    })
    print("  [OK] Archetype 05: KPI Metrics rendered")

    # 6. Slide 6: Feature Matrix / 3 Pillars
    s6 = prs.slides.add_slide(blank_layout)
    render_feature_list(s6, theme, {
        "category": "CORE CAPABILITIES",
        "title": "系统三大核心支柱与特性矩阵",
        "subtitle": "打通色彩感知算法、表现力空间几何与原生矢量编译的技术闭环",
        "features": [
            {
                "tag": "PILLAR 01", "icon": "auto_awesome",
                "title": "HCT 动态感知色彩", "subtitle": "Dual-Chroma Dynamic Palette",
                "desc": "基于 Google MCU 算法推导和谐高对比度的 Primary 与 Tertiary 双强调色体系，确保色彩层级清晰鲜活。",
                "points": ["WCAG AAA 级无障碍对比度", "Light/Dark 双模式毫秒级映射", "4 套旗舰级预设调色板开箱即用"],
                "badge": "官方算法内核", "highlight": False
            },
            {
                "tag": "PILLAR 02", "icon": "grid_view",
                "title": "Bento 网格与大圆角", "subtitle": "28dp Modular Containers",
                "desc": "基于 Google Design Tokens 标准，采用 28dp 核心大容器倒角与 9999px 全胶囊药丸交互组件。",
                "points": ["DrawingML 几何曲率高保真控制", "模块化信息分舱与呼吸感留白", "多端自适应安全内边距"],
                "badge": "结构美学重构", "highlight": True
            },
            {
                "tag": "PILLAR 03", "icon": "verified",
                "title": "原生矢量与抗降级", "subtitle": "Pure DrawingML Rendering",
                "desc": "100% 原生矢量图层编译。中文排版注入双字族声明，彻底杜绝回退宋体问题，跨平台完全可二次编辑。",
                "points": ["33+ 款官方 Material Symbols 矢量", "跨设备字体渲染永不走样", "支持 PowerPoint/Keynote 自由拖拽编辑"],
                "badge": "100% 原生可编辑", "highlight": False
            }
        ],
        "footer_right": "Page 06 / 08"
    })
    print("  [OK] Archetype 06: Feature Matrix rendered")

    # 7. Slide 7: Quote & Key Takeaway
    s7 = prs.slides.add_slide(blank_layout)
    render_quote_takeaway(s7, theme, {
        "category": "EXECUTIVE TAKEAWAY",
        "title": "核心观点与设计哲学共识",
        "subtitle": "下一代演示体验的底层逻辑：秩序、活力与无损表达",
        "quote_text": "“设计不仅是视觉的外表，更是信息传达的秩序。Material 3 Expressive 将感知色彩与结构表现力融合，让每一个关键决策洞察都获得最清晰的共鸣。”",
        "quote_author": "Google Material Design Core Team",
        "quote_role": "Design Principles & Systems Architecture",
        "pillars": [
            {
                "icon": "auto_awesome", "title": "高对比度聚焦",
                "desc": "以 Primary + Tertiary 形成双视觉锚点，彻底杜绝单调冷灰对受众注意力的稀释。"
            },
            {
                "icon": "grid_view", "title": "便当盒模块化分舱",
                "desc": "采用 28dp 容器大倒角与呼吸感间距，建立严密清晰的信息层级骨架。"
            },
            {
                "icon": "verified", "title": "100% 原生矢量可编辑",
                "desc": "基于 DrawingML 纯净编译，跨平台自由缩放编辑，文字永不发生降级走样。"
            }
        ],
        "footer_right": "Page 07 / 08"
    })
    print("  [OK] Archetype 07: Quote Takeaway rendered")

    # 8. Slide 8: Team & Talent Showcase
    s8 = prs.slides.add_slide(blank_layout)
    render_team_showcase(s8, theme, {
        "category": "ORGANIZATION & TALENT",
        "title": "核心专家团队与研发组织架构",
        "subtitle": "跨学科专家团队共同驱动底层算法、渲染内核与设计规范的深度融合",
        "members": [
            {
                "name": "Dr. Elena Vance", "role": "CHIEF ARCHITECT", "icon": "psychology",
                "desc": "Google DeepMind 先进编码团队专家，主导 M3 Expressive HCT 感知色彩与 DrawingML 动态编译引擎。",
                "tags": ["HCT 算法", "DrawingML", "管线架构", "动态着色"]
            },
            {
                "name": "Marcus Zhang", "role": "DESIGN SYSTEMS LEAD", "icon": "palette",
                "desc": "10 年前沿 UI/UX 系统专家，负责 28dp Bento 容器美学、有机异形花瓣与排版规范设计。",
                "tags": ["Bento 架构", "M3 Tokens", "无衬线排版", "几何倒角"]
            },
            {
                "name": "Sarah Lin", "role": "ENGINEERING PRINCIPAL", "icon": "account_tree",
                "desc": "专注于高性能跨平台文档编译、矢量图标动态重着色管线与并发安全保护系统构建。",
                "tags": ["Rust 渲染", "矢量引擎", "工业化落地", "跨端编译"]
            }
        ],
        "footer_right": "Page 08 / 08"
    })
    print("  [OK] Archetype 08: Team Showcase rendered")

    # Save complete deck
    out_file = safe_save_pptx(prs, "tests/output/m3_archetypes_showcase.pptx")
    print(f"\n[PASS] Successfully generated 8-slide presentation: {out_file}")
    return out_file

if __name__ == "__main__":
    build_showcase_deck()
