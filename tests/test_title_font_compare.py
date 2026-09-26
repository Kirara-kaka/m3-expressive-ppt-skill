from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Test A: <a:ea> omitted, only <a:latin typeface="Segoe UI Variable Display"/>
tbA = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(1.2))
pA = tbA.text_frame.paragraphs[0]
rA = pA.add_run()
rA.text = "方案A：只设 latin 为 Segoe UI Variable Display（不写 ea 标签）"
rA.font.size = Pt(24)
rA.font.bold = True
rPrA = rA._r.get_or_add_rPr()
rPrA.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))

# Test B: <a:latin typeface="Segoe UI Variable Display"/> + <a:ea typeface="Noto Sans SC"/>
tbB = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.2))
pB = tbB.text_frame.paragraphs[0]
rB = pB.add_run()
rB.text = "方案B：latin 为 Segoe UI Variable Display + ea 为 Noto Sans SC (Bold)"
rB.font.size = Pt(24)
rB.font.bold = True
rPrB = rB._r.get_or_add_rPr()
rPrB.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))
rPrB.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Noto Sans SC"/>'))

# Test C: <a:latin typeface="Segoe UI Variable Display"/> + <a:ea typeface="Microsoft YaHei UI"/>
tbC = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11), Inches(1.2))
pC = tbC.text_frame.paragraphs[0]
rC = pC.add_run()
rC.text = "方案C：latin 为 Segoe UI Variable Display + ea 为 Microsoft YaHei UI (Bold)"
rC.font.size = Pt(24)
rC.font.bold = True
rPrC = rC._r.get_or_add_rPr()
rPrC.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))
rPrC.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Microsoft YaHei UI"/>'))

# Test D: Exactly user's instruction: <a:ea typeface="Segoe UI Variable Display"/>
tbD = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11), Inches(1.2))
pD = tbD.text_frame.paragraphs[0]
rD = pD.add_run()
rD.text = "方案D：用户指令 ea 强设为 Segoe UI Variable Display（发生宋体降级）"
rD.font.size = Pt(24)
rD.font.bold = True
rPrD = rD._r.get_or_add_rPr()
rPrD.append(parse_xml('<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))
rPrD.append(parse_xml('<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Segoe UI Variable Display"/>'))

prs.save("tests/output/test_title_font_compare.pptx")
print("Saved test_title_font_compare.pptx")
