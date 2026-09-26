import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

svg_content = '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100" viewBox="0 0 24 24"><circle cx="12" cy="12" r="10" fill="#6750A4"/></svg>'
with open('test_icon.svg', 'w') as f:
    f.write(svg_content)

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

# 1. Test shape adjustment for roundRect
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(1))
prstGeom = shape._element.spPr.prstGeom
prstGeom.remove(prstGeom.avLst)
new_av = parse_xml('<a:avLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:gd name="adj" fmla="val 50000"/></a:avLst>')
prstGeom.append(new_av)
print("Pill adjustment configured!")

# 2. Test SVG insertion
try:
    slide.shapes.add_picture('test_icon.svg', Inches(1), Inches(3), Inches(1), Inches(1))
    print("Native SVG add_picture succeeded!")
except Exception as e:
    print("Native SVG failed:", type(e), e)

prs.save('test_output.pptx')
print("Saved test_output.pptx successfully!")
