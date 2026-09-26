"""
Google Material 3 Expressive Component Library for PowerPoint.
Provides M3Card, M3Chip, M3HeroMetric, and M3BentoGrid layout generators.
"""

import os
from dataclasses import dataclass
from typing import Tuple, List, Optional
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

from core.tokens import M3Theme
from core.shapes_tokens import M3_CORNER_EXTRA_LARGE, M3_CORNER_FULL
from core.icon_manager import get_recolored_icon_png
from engine.geometry import apply_dp_corner, apply_pill_corner, apply_ambient_shadow
from engine.typography import (
    add_styled_paragraph, 
    create_textbox,
    SCALE_DISPLAY_LARGE,
    SCALE_HEADLINE,
    SCALE_TITLE,
    SCALE_BODY,
    SCALE_BODY_SMALL,
    SCALE_LABEL
)

@dataclass
class CardBounds:
    left: float
    top: float
    width: float
    height: float

@dataclass
class CardResult:
    shape: any
    bounds: CardBounds
    inner_bounds: CardBounds

def remove_shape_shadow(shape):
    """
    Remove default PowerPoint theme drop shadow (idx='2' in effectRef) from a shape.
    Ensures crisp, flat Material 3 components without muddy downward blur.
    """
    style = shape._sp.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        effectRef = style.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
        if effectRef is not None:
            effectRef.set('idx', '0')

class M3Components:
    """Component builder targeting Google Material 3 Expressive presentation standards."""

    @staticmethod
    def create_card(slide,
                    theme: M3Theme,
                    left: float,
                    top: float,
                    width: float,
                    height: float,
                    variant: str = "filled",
                    color_role: str = "surface_container",
                    corner_dp: int = M3_CORNER_EXTRA_LARGE,
                    padding: float = 0.28) -> CardResult:
        """
        Create a Material 3 Expressive Card Container.
        
        Variants:
          - 'filled': Solid container color, no outline.
          - 'elevated': 'surface_container_lowest' with soft tinted ambient shadow.
          - 'outlined': Subtle 0.75pt 'outline_variant' border.
        """
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        apply_dp_corner(shape, width, height, corner_dp=corner_dp)

        if variant == "elevated":
            shape.fill.solid()
            shape.fill.fore_color.rgb = theme.rgb("surface_container_lowest")
            shape.line.fill.background()
            apply_ambient_shadow(shape, hex_color=theme.hex("primary"), alpha_val="8000")
        elif variant == "outlined":
            shape.fill.solid()
            shape.fill.fore_color.rgb = theme.rgb(color_role)
            shape.line.color.rgb = theme.rgb("outline_variant")
            shape.line.width = Pt(0.75)
            remove_shape_shadow(shape)
        else: # 'filled'
            shape.fill.solid()
            shape.fill.fore_color.rgb = theme.rgb(color_role)
            shape.line.fill.background()
            remove_shape_shadow(shape)

        bounds = CardBounds(left, top, width, height)
        inner_bounds = CardBounds(
            left=left + Inches(padding),
            top=top + Inches(padding),
            width=width - Inches(padding * 2),
            height=height - Inches(padding * 2)
        )
        return CardResult(shape=shape, bounds=bounds, inner_bounds=inner_bounds)

    @staticmethod
    def create_chip(slide,
                    theme: M3Theme,
                    left: float,
                    top: float,
                    text: str,
                    icon_name: Optional[str] = None,
                    color_role: str = "primary_container",
                    on_color_role: str = "on_primary_container",
                    height: float = 0.38,
                    width: Optional[float] = None) -> any:
        """
        Create a 100% capsule Pill Chip with optional Material Symbol icon.
        If width is provided (in inches), uses exact width; otherwise calculates dynamically from text length.
        """
        has_icon = icon_name is not None
        zh_count = sum(1 for c in text if '\u4e00' <= c <= '\u9fff')
        en_count = len(text) - zh_count
        
        if width is not None:
            w_emu = Inches(width) if isinstance(width, (int, float)) and width < 100 else width
        else:
            text_width = (en_count * 0.096) + (zh_count * 0.155)
            w_emu = Inches(0.36 + text_width + (0.30 if has_icon else 0))
        
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w_emu, Inches(height))
        pill.fill.solid()
        pill.fill.fore_color.rgb = theme.rgb(color_role)
        pill.line.fill.background()
        apply_pill_corner(pill)
        remove_shape_shadow(pill)

        if has_icon:
            icon_png = get_recolored_icon_png(icon_name, theme.hex(on_color_role), size=128)
            icon_size = Inches(height * 0.58)
            icon_top = top + Inches((height - (height * 0.58)) / 2)
            slide.shapes.add_picture(icon_png, left + Inches(0.12), icon_top, icon_size, icon_size)
            
            # Text box with icon offset, vertically centered
            tb = create_textbox(
                slide, 
                left=left + Inches(0.12 + (height * 0.58) + 0.08), 
                top=top, 
                width=w_emu - Inches(0.12 + (height * 0.58) + 0.10), 
                height=Inches(height),
                vertical_anchor=MSO_ANCHOR.MIDDLE
            )
            tb.text_frame.word_wrap = False
            add_styled_paragraph(tb.text_frame, text, size_pt=SCALE_LABEL, bold=True, 
                                color_rgb=theme.rgb(on_color_role), align=PP_ALIGN.LEFT)
        else:
            tb = create_textbox(
                slide, 
                left=left, 
                top=top, 
                width=w_emu, 
                height=Inches(height),
                vertical_anchor=MSO_ANCHOR.MIDDLE
            )
            tb.text_frame.word_wrap = False
            add_styled_paragraph(tb.text_frame, text, size_pt=SCALE_LABEL, bold=True, 
                                color_rgb=theme.rgb(on_color_role), align=PP_ALIGN.CENTER)
            
        return pill

    @staticmethod
    def create_hero_metric(slide,
                           theme: M3Theme,
                           left: float,
                           top: float,
                           width: float,
                           height: float,
                           metric_value: str,
                           metric_label: str,
                           trend_str: Optional[str] = None,
                           color_role: str = "primary",
                           bg_role: str = "surface_container_lowest",
                           label_color_role: Optional[str] = None) -> any:
        """
        Create a high-impact KPI Hero Metric Card with generous whitespace and zero overlap.
        """
        card_res = M3Components.create_card(
            slide, theme, left, top, width, height, 
            variant="elevated" if bg_role == "surface_container_lowest" else "filled",
            color_role=bg_role, corner_dp=24, padding=0.22
        )
        
        # Calculate trend chip dimensions and layout
        chip_w_in = 0.0
        if trend_str:
            is_positive = "+" in trend_str or "↑" in trend_str
            if bg_role in ["primary", "secondary", "tertiary", "inverse_surface"]:
                chip_color = "tertiary_container" if is_positive else "surface_bright"
                chip_on_color = "on_tertiary_container" if is_positive else "on_surface"
            else:
                chip_color = "tertiary_container" if is_positive else "surface_container_high"
                chip_on_color = "on_tertiary_container" if is_positive else "on_surface_variant"
            
            zh_c = sum(1 for c in trend_str if '\u4e00' <= c <= '\u9fff')
            en_c = len(trend_str) - zh_c
            chip_w_in = 0.36 + (en_c * 0.096) + (zh_c * 0.155) + 0.30

            if card_res.inner_bounds.width >= Inches(2.8):
                # Modern Dashboard KPI: Place trend chip in top-right corner
                chip_left = card_res.inner_bounds.left + card_res.inner_bounds.width - Inches(chip_w_in)
                chip_top = card_res.inner_bounds.top + Inches(0.06)
                M3Components.create_chip(
                    slide, theme,
                    chip_left, chip_top,
                    trend_str,
                    icon_name="growth" if is_positive else "trending",
                    color_role=chip_color,
                    on_color_role=chip_on_color,
                    height=0.34
                )
                tb_w = card_res.inner_bounds.width - Inches(chip_w_in + 0.15)
                tb_h = card_res.inner_bounds.height
            else:
                # Narrow container: place at bottom inside bounds
                chip_left = card_res.inner_bounds.left
                chip_top = card_res.inner_bounds.top + card_res.inner_bounds.height - Inches(0.36)
                M3Components.create_chip(
                    slide, theme,
                    chip_left, chip_top,
                    trend_str,
                    icon_name="growth" if is_positive else "trending",
                    color_role=chip_color,
                    on_color_role=chip_on_color,
                    height=0.32
                )
        else:
            tb_w = card_res.inner_bounds.width
            tb_h = card_res.inner_bounds.height

        # Textbox dimensions
        tb_w = card_res.inner_bounds.width - (Inches(chip_w_in + 0.15) if (trend_str and card_res.inner_bounds.width >= Inches(2.8)) else 0)
        
        # 1. Big Metric Number (top 0.50 in)
        tb_num = create_textbox(slide, card_res.inner_bounds.left, card_res.inner_bounds.top, tb_w, Inches(0.50))
        add_styled_paragraph(tb_num.text_frame, metric_value, size_pt=SCALE_DISPLAY_LARGE, bold=True, 
                             color_rgb=theme.rgb(color_role))
        
        # Determine contrast-compliant label color
        if label_color_role is None:
            if bg_role in ["primary", "secondary", "tertiary", "inverse_surface"] or (color_role and color_role.startswith("on_")):
                if bg_role == "primary" or color_role == "on_primary":
                    label_color_role = "on_primary"
                elif bg_role == "secondary" or color_role == "on_secondary":
                    label_color_role = "on_secondary"
                elif bg_role == "tertiary" or color_role == "on_tertiary":
                    label_color_role = "on_tertiary"
                elif bg_role == "inverse_surface":
                    label_color_role = "inverse_on_surface"
                else:
                    label_color_role = color_role
            elif bg_role == "primary_container":
                label_color_role = "on_primary_container"
            elif bg_role == "tertiary_container":
                label_color_role = "on_tertiary_container"
            elif bg_role == "secondary_container":
                label_color_role = "on_secondary_container"
            else:
                label_color_role = "on_surface_variant"

        # 2. Metric Label (under number, generous spacing to prevent overlap)
        tb_lbl = create_textbox(slide, card_res.inner_bounds.left, card_res.inner_bounds.top + Inches(0.56), tb_w, Inches(0.38))
        add_styled_paragraph(tb_lbl.text_frame, metric_label, size_pt=SCALE_BODY_SMALL, bold=False, 
                             color_rgb=theme.rgb(label_color_role))

        return card_res

    @staticmethod
    def create_icon_container(slide,
                              theme: M3Theme,
                              left: float,
                              top: float,
                              size: float = 0.44,
                              icon_name: str = "star",
                              color_role: str = "primary_container",
                              icon_color_role: str = "primary",
                              shape_type: str = "circle") -> any:
        """
        Create an M3 Expressive Tonal Icon Container (circle, squircle or rounded rect).
        Encases raw icons into a branded tactile badge.
        """
        l_emu = Inches(left) if isinstance(left, (int, float)) and left < 100 else left
        t_emu = Inches(top) if isinstance(top, (int, float)) and top < 100 else top
        s_emu = Inches(size) if isinstance(size, (int, float)) and size < 100 else size
        s_in = size if (isinstance(size, (int, float)) and size < 100) else (size / 914400.0)

        mso_shape = MSO_SHAPE.OVAL if shape_type == "circle" else MSO_SHAPE.ROUNDED_RECTANGLE
        badge = slide.shapes.add_shape(mso_shape, l_emu, t_emu, s_emu, s_emu)
        badge.fill.solid()
        badge.fill.fore_color.rgb = theme.rgb(color_role)
        badge.line.fill.background()
        remove_shape_shadow(badge)
        if shape_type == "squircle":
            apply_pill_corner(badge)
        elif shape_type == "rounded":
            apply_dp_corner(badge, 16)
            
        icon_inner_size = Inches(s_in * 0.58)
        offset = Inches((s_in - (s_in * 0.58)) / 2.0)
        icon_png = get_recolored_icon_png(icon_name, theme.hex(icon_color_role), size=128)
        slide.shapes.add_picture(icon_png, l_emu + offset, t_emu + offset, icon_inner_size, icon_inner_size)
        return badge

    @staticmethod
    def create_progress_bar(slide,
                            theme: M3Theme,
                            left: any,
                            top: any,
                            width: any,
                            height: any = 0.08,
                            progress: float = 0.75,
                            track_color_role: str = "surface_container_highest",
                            indicator_color_role: str = "primary") -> tuple:
        """
        Create a Google M3 100% capsule Linear Progress Indicator.
        Both track and active indicator have full pill rounded caps.
        """
        l_emu = Inches(left) if isinstance(left, (int, float)) and left < 100 else left
        t_emu = Inches(top) if isinstance(top, (int, float)) and top < 100 else top
        w_emu = Inches(width) if isinstance(width, (int, float)) and width < 100 else width
        h_emu = Inches(height) if isinstance(height, (int, float)) and height < 100 else height

        # 1. Background Track
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_emu, t_emu, w_emu, h_emu)
        track.fill.solid()
        track.fill.fore_color.rgb = theme.rgb(track_color_role)
        track.line.fill.background()
        apply_pill_corner(track)
        remove_shape_shadow(track)
        
        # 2. Foreground Active Indicator
        clamped_prog = max(0.01, min(1.0, progress))
        ind_w = max(w_emu * clamped_prog, h_emu)
        indicator = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_emu, t_emu, ind_w, h_emu)
        indicator.fill.solid()
        indicator.fill.fore_color.rgb = theme.rgb(indicator_color_role)
        indicator.line.fill.background()
        apply_pill_corner(indicator)
        remove_shape_shadow(indicator)
        
        return (track, indicator)

    @staticmethod
    def create_checklist(slide,
                         theme: M3Theme,
                         left: float,
                         top: float,
                         width: float,
                         items: list,
                         icon_name: str = "check",
                         icon_bg_role: str = "primary_container",
                         icon_fg_role: str = "primary",
                         font_size_pt: float = SCALE_BODY_SMALL,
                         text_color_role: str = "on_surface",
                         bold_text: bool = False,
                         item_spacing_in: float = 0.36,
                         badge_size: float = 0.22) -> float:
        """
        Render a modern M3 checklist where each item has an aligned micro-icon badge
        instead of a plain text bullet. Returns the total height occupied.
        Both the micro-badge and the text box are strictly vertically centered along the exact same horizontal axis.
        """
        cur_y = Inches(top) if isinstance(top, (int, float)) and top < 100 else top
        l_emu = Inches(left) if isinstance(left, (int, float)) and left < 100 else left
        w_emu = Inches(width) if isinstance(width, (int, float)) and width < 100 else width
        b_size_emu = Inches(badge_size) if isinstance(badge_size, (int, float)) and badge_size < 100 else badge_size
        spacing_emu = Inches(item_spacing_in) if isinstance(item_spacing_in, (int, float)) and item_spacing_in < 100 else item_spacing_in

        # Optical offset for single-line East Asian text in PowerPoint:
        # At zero-margins, the optical midline of East Asian characters lands at ~0.72 * font_size below textbox top.
        tb_offset_emu = Pt(font_size_pt * 0.72)

        for item in items:
            if isinstance(item, dict):
                text = item.get("text", "")
                item_icon = item.get("icon", icon_name)
                item_bg = item.get("bg_role", icon_bg_role)
                item_fg = item.get("fg_role", icon_fg_role)
            else:
                text = str(item)
                item_icon = icon_name
                item_bg = icon_bg_role
                item_fg = icon_fg_role

            row_mid_y = cur_y + (b_size_emu // 2)

            # Micro circle badge with icon - centered at row_mid_y
            M3Components.create_icon_container(
                slide, theme,
                left=l_emu, 
                top=cur_y,
                size=badge_size,
                icon_name=item_icon,
                color_role=item_bg,
                icon_color_role=item_fg,
                shape_type="circle"
            )

            # Aligned text box - optical midline lands exactly on row_mid_y
            text_offset_in = badge_size + 0.10
            text_offset_emu = Inches(text_offset_in)
            text_left = l_emu + text_offset_emu
            text_w = w_emu - text_offset_emu
            
            tb = create_textbox(
                slide, 
                left=text_left, 
                top=row_mid_y - tb_offset_emu, 
                width=text_w, 
                height=Inches(0.25),
                margin_zero=True
            )
            add_styled_paragraph(tb.text_frame, text, size_pt=font_size_pt, bold=bold_text,
                                 color_rgb=theme.rgb(text_color_role))
            cur_y += spacing_emu

        top_emu = Inches(top) if isinstance(top, (int, float)) and top < 100 else top
        return (cur_y - top_emu) / 914400.0

class M3Canvas:
    """Manages full-slide canvas background, dynamic tinting, and ambient blooms."""

    @staticmethod
    def apply_background(slide, 
                         theme: M3Theme, 
                         style: str = "tinted", 
                         add_ambient_bloom: bool = True) -> any:
        """
        Apply a theme-unified background to the slide canvas.
        
        Styles:
          - 'tinted': Uses 'surface_container_low' for distinct brand tone (not plain white).
          - 'surface': Uses subtle 'surface' (tone 98).
          - 'container': Uses richer 'surface_container' (tone 94) for high card contrast.
        """
        # 1. Full-bleed background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        
        if style == "container":
            bg.fill.fore_color.rgb = theme.rgb("surface_container")
        elif style == "surface":
            bg.fill.fore_color.rgb = theme.rgb("surface")
        else: # 'tinted' (default)
            bg.fill.fore_color.rgb = theme.rgb("surface_container_low")
            
        bg.line.fill.background()

        # 2. Expressive Ambient Bloom (Soft glowing organic halo in top-right or corner)
        if add_ambient_bloom:
            bloom_path = os.path.join(BASE_DIR, "assets", "icons", "rendered", f"bloom_{theme.seed_hex.lstrip('#')}.png")
            if not os.path.exists(bloom_path):
                M3Canvas._create_ambient_bloom_png(bloom_path, theme.rgb_tuple("tertiary_container"), theme.rgb_tuple("primary_container"))
            
            # Place subtle bloom in the top-right corner behind cards
            bloom_size = Inches(5.5)
            slide.shapes.add_picture(bloom_path, Inches(8.5), Inches(-1.0), bloom_size, bloom_size)

        return bg

    @staticmethod
    def _create_ambient_bloom_png(output_path: str, color1_rgb: tuple, color2_rgb: tuple, size: int = 512):
        """Create a soft radial ambient bloom PNG with Gaussian falloff."""
        from PIL import Image, ImageDraw
        img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        cx, cy = size / 2, size / 2
        max_r = size * 0.48

        # Draw concentric soft circles with decaying alpha
        r_step = 6
        for r in range(int(max_r), 0, -r_step):
            # Alpha decreases from center out
            factor = (1.0 - (r / max_r)) ** 2
            alpha = int(45 * factor) # soft max opacity ~18%
            # Blend color1 (tertiary) and color2 (primary)
            t = r / max_r
            cr = int(color1_rgb[0] * (1 - t) + color2_rgb[0] * t)
            cg = int(color1_rgb[1] * (1 - t) + color2_rgb[1] * t)
            cb = int(color1_rgb[2] * (1 - t) + color2_rgb[2] * t)
            draw.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(cr, cg, cb, alpha))

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        img.save(output_path, "PNG")


class M3BentoGrid:
    """Mathematical Grid Partitioning for Material 3 Bento Layouts."""

    @staticmethod
    def layout_1_hero_3_sub(x: float, y: float, w: float, h: float, gap: float = 0.22) -> Tuple[CardBounds, CardBounds, CardBounds, CardBounds]:
        """
        Layout:
        [   Hero Card (Left)   ] [ Sub Card Top (Right)         ]
        [                      ] [ Sub Card BL  ] [ Sub Card BR  ]
        """
        hero_w = (w - gap) * 0.48
        sub_w = w - hero_w - gap
        
        # Left Hero Card
        b_hero = CardBounds(x, y, hero_w, h)
        
        # Right Top Card
        top_h = (h - gap) * 0.48
        b_top = CardBounds(x + hero_w + gap, y, sub_w, top_h)
        
        # Right Bottom Cards (split horizontally into 2 mini cards)
        bot_y = y + top_h + gap
        bot_h = h - top_h - gap
        bot_mini_w = (sub_w - gap) / 2.0
        
        b_bot_left = CardBounds(x + hero_w + gap, bot_y, bot_mini_w, bot_h)
        b_bot_right = CardBounds(x + hero_w + gap + bot_mini_w + gap, bot_y, bot_mini_w, bot_h)
        
        return b_hero, b_top, b_bot_left, b_bot_right

    @staticmethod
    def layout_3_columns(x: float, y: float, w: float, h: float, gap: float = 0.22) -> List[CardBounds]:
        """3 equal-width vertical columns."""
        col_w = (w - (gap * 2)) / 3.0
        return [
            CardBounds(x + i * (col_w + gap), y, col_w, h)
            for i in range(3)
        ]

    @staticmethod
    def layout_2x2(x: float, y: float, w: float, h: float, gap: float = 0.22) -> List[CardBounds]:
        """2x2 balanced grid."""
        cell_w = (w - gap) / 2.0
        cell_h = (h - gap) / 2.0
        return [
            CardBounds(x, y, cell_w, cell_h),
            CardBounds(x + cell_w + gap, y, cell_w, cell_h),
            CardBounds(x, y + cell_h + gap, cell_w, cell_h),
            CardBounds(x + cell_w + gap, y + cell_h + gap, cell_w, cell_h),
        ]
