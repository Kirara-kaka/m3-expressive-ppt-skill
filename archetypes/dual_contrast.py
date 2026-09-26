"""
Archetype 04: Dual-Contrast Comparison Slide.
Side-by-side comparison matrix leveraging M3 Primary vs Tertiary dual-chroma pairing.
"""

from pptx.util import Inches
from core.tokens import M3Theme
from core.icon_manager import get_recolored_icon_png
from engine.components import M3Components
from engine.typography import create_textbox, add_styled_paragraph, SCALE_HEADLINE, SCALE_TITLE, SCALE_BODY, SCALE_BODY_SMALL
from archetypes.base import apply_slide_canvas, render_standard_header, render_standard_footer

def render_dual_contrast(slide, theme: M3Theme, data: dict = None):
    """
    Render Dual-Contrast Comparison Matrix.
    
    Data Schema:
      - category, title, subtitle
      - left_plan:
          - chip: "传统单色方案"
          - title: "传统沉闷低对比度设计"
          - desc: "..."
          - points: ["...", "..."]
          - metric_val: "34.5%"
          - metric_label: "阅读注意留存率"
      - right_plan:
          - chip: "M3 EXPRESSIVE 演进方案"
          - title: "高活力双强调色感知架构"
          - desc: "..."
          - points: ["...", "..."]
          - metric_val: "86.4%"
          - metric_label: "视觉聚焦点转换率"
    """
    if data is None:
        data = {}

    apply_slide_canvas(slide, theme, style="tinted", add_ambient_bloom=True)

    render_standard_header(
        slide, theme,
        category_text=data.get("category", "ARCHITECTURE COMPARISON"),
        title_text=data.get("title", "传统演示方案 vs M3 Expressive 深度比对"),
        subtitle_text=data.get("subtitle", "对比传统沉闷设计局限与下一代表现力架构的体验跃迁"),
        category_icon=data.get("category_icon", "compare_arrows")
    )

    card_y = 2.05
    card_h = 4.7
    total_w = 11.733
    gap = 0.36
    card_w = (total_w - gap) / 2.0
    left_x = 0.8
    right_x = left_x + card_w + gap

    # -------------------------------------------------------------
    # PLAN DATA EXTRACTION
    # -------------------------------------------------------------
    l_data = data.get("left_plan", {
        "chip": "TRADITIONAL BASELINE",
        "title": "传统低对比度设计局限",
        "desc": "单一冷灰缺乏视觉聚焦，受众阅读易产生认知疲劳。",
        "points": [
            "单调沉闷的四角直角与表格排版",
            "黑重脏阴影破坏画面透气呼吸感",
            "中文字体在跨端时易降级退化为宋体"
        ],
        "metric_val": "34.5%",
        "metric_label": "信息瞬时记忆留存率",
        "trend": "传统基准参考"
    })

    r_data = data.get("right_plan", {
        "chip": "M3 EXPRESSIVE EVOLVED",
        "title": "高活力双强调色感知架构",
        "desc": "基于 HCT 算法引入双强调色，重塑高对比度视觉层次。",
        "points": [
            "28dp 核心大圆角与 9999px 全胶囊药丸",
            "Surface Container 纯净色阶代替脏黑阴影",
            "DrawingML 双字族注入，现代几何无衬线抗降级"
        ],
        "metric_val": "86.4%",
        "metric_label": "核心决策信息传达与留存率",
        "trend": "+2.5x 跃迁"
    })

    l_points = l_data.get("points", [])
    r_points = r_data.get("points", [])
    max_pts = max(len(l_points), len(r_points), 3)

    item_spacing_in = 0.36 if max_pts <= 3 else 0.30
    tb_h_in = 0.24
    total_checklist_h = (max_pts - 1) * item_spacing_in + tb_h_in
    sub_card_h = 1.42 if max_pts <= 3 else 1.62
    top_pad_in = (sub_card_h - total_checklist_h) / 2.0
    metric_h = 1.18

    # -------------------------------------------------------------
    # LEFT CARD: Baseline (Elevated White Card)
    # -------------------------------------------------------------
    card_l = M3Components.create_card(
        slide, theme,
        left=Inches(left_x), top=Inches(card_y),
        width=Inches(card_w), height=Inches(card_h),
        variant="elevated", color_role="surface_container_lowest", corner_dp=28, padding=0.28
    )
    inner_l = card_l.inner_bounds
    sub_card_y = inner_l.top + Inches(1.36)
    metric_top = inner_l.top + Inches(2.96)

    # Top Chip
    M3Components.create_chip(
        slide, theme, left=inner_l.left, top=inner_l.top,
        text=l_data.get("chip", "TRADITIONAL BASELINE"), icon_name="tune",
        color_role="surface_container_high", on_color_role="on_surface", height=0.32
    )

    # Title & Concise Desc
    tb_lt = create_textbox(slide, inner_l.left, inner_l.top + Inches(0.46), inner_l.width, Inches(0.85))
    add_styled_paragraph(tb_lt.text_frame, l_data.get("title", "传统低对比度设计局限"), 
                         size_pt=SCALE_HEADLINE, bold=True, color_rgb=theme.rgb("on_surface"))
    add_styled_paragraph(tb_lt.text_frame, l_data.get("desc", ""), 
                         size_pt=SCALE_BODY_SMALL, bold=False, color_rgb=theme.rgb("on_surface_variant"), space_before_pt=2)

    # Sub-card Container wrapping Checklist Points (surface_container_high)
    sub_l = M3Components.create_card(
        slide, theme,
        left=inner_l.left, top=sub_card_y,
        width=inner_l.width, height=Inches(sub_card_h),
        variant="filled", color_role="surface_container_high", corner_dp=16, padding=0.0
    )

    # Exact vertically centered checklist (symmetrical top & bottom padding)
    M3Components.create_checklist(
        slide, theme,
        left=inner_l.left + Inches(0.18),
        top=sub_card_y + Inches(top_pad_in),
        width=inner_l.width - Inches(0.36),
        items=l_points,
        icon_name="close",
        icon_bg_role="surface_container_lowest",
        icon_fg_role="outline",
        font_size_pt=10.5,
        text_color_role="on_surface_variant",
        item_spacing_in=item_spacing_in,
        badge_size=0.18
    )

    # Metric Box (surface_container_low, distinct from wrapping sub-card)
    M3Components.create_hero_metric(
        slide, theme,
        left=inner_l.left, top=metric_top,
        width=inner_l.width, height=Inches(metric_h),
        metric_value=l_data.get("metric_val", "34.5%"),
        metric_label=l_data.get("metric_label", "信息瞬时记忆留存率"),
        trend_str=l_data.get("trend", "传统基准参考"),
        color_role="on_surface", bg_role="surface_container_low"
    )

    # -------------------------------------------------------------
    # RIGHT CARD: M3 Expressive Evolved (Primary Container)
    # -------------------------------------------------------------
    card_r = M3Components.create_card(
        slide, theme,
        left=Inches(right_x), top=Inches(card_y),
        width=Inches(card_w), height=Inches(card_h),
        variant="filled", color_role="primary_container", corner_dp=28, padding=0.28
    )
    inner_r = card_r.inner_bounds

    M3Components.create_chip(
        slide, theme, left=inner_r.left, top=inner_r.top,
        text=r_data.get("chip", "M3 EXPRESSIVE EVOLVED"), icon_name="verified",
        color_role="primary", on_color_role="on_primary", height=0.32
    )

    tb_rt = create_textbox(slide, inner_r.left, inner_r.top + Inches(0.46), inner_r.width, Inches(0.85))
    add_styled_paragraph(tb_rt.text_frame, r_data.get("title", "高活力双强调色感知架构"), 
                         size_pt=SCALE_HEADLINE, bold=True, color_rgb=theme.rgb("on_primary_container"))
    add_styled_paragraph(tb_rt.text_frame, r_data.get("desc", ""), 
                         size_pt=SCALE_BODY_SMALL, bold=False, color_rgb=theme.rgb("on_primary_container"), space_before_pt=2)

    # Sub-card Container wrapping Checklist Points (Pure White Container)
    sub_r = M3Components.create_card(
        slide, theme,
        left=inner_r.left, top=sub_card_y,
        width=inner_r.width, height=Inches(sub_card_h),
        variant="filled", color_role="surface_container_lowest", corner_dp=16, padding=0.0
    )

    # Exact vertically centered checklist (symmetrical top & bottom padding)
    M3Components.create_checklist(
        slide, theme,
        left=inner_r.left + Inches(0.18),
        top=sub_card_y + Inches(top_pad_in),
        width=inner_r.width - Inches(0.36),
        items=r_points,
        icon_name="check",
        icon_bg_role="primary",
        icon_fg_role="on_primary",
        font_size_pt=10.5,
        text_color_role="primary",
        bold_text=True,
        item_spacing_in=item_spacing_in,
        badge_size=0.18
    )

    # Metric Box (Solid Primary Hero Container, distinct from pure white wrapping sub-card)
    M3Components.create_hero_metric(
        slide, theme,
        left=inner_r.left, top=metric_top,
        width=inner_r.width, height=Inches(metric_h),
        metric_value=r_data.get("metric_val", "86.4%"),
        metric_label=r_data.get("metric_label", "关键指标提升"),
        trend_str=r_data.get("trend", "+2.5x 跃迁"),
        color_role="on_primary", bg_role="primary",
        label_color_role="on_primary"
    )

    # -------------------------------------------------------------
    # CENTER "VS" BADGE (Drawn last to float on top of both cards)
    # -------------------------------------------------------------
    center_mid_x = left_x + card_w + (gap / 2.0)
    vs_y = card_y + (card_h / 2.0) - 0.22
    M3Components.create_chip(
        slide, theme,
        left=Inches(center_mid_x - 0.38), top=Inches(vs_y),
        text="VS", icon_name="bolt",
        color_role="tertiary", on_color_role="on_tertiary", height=0.44
    )

    render_standard_footer(slide, theme, left_meta=data.get("footer_left", "Dual-Contrast Comparison · M3 Expressive"), right_meta=data.get("footer_right", ""))
