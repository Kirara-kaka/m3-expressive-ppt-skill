"""
Archetype 08: Team & Subject Showcase Slide.
Persona / Team / Product portfolio cards featuring avatar containers, role pills, and expertise tags.
"""

from pptx.util import Inches
from core.tokens import M3Theme
from core.icon_manager import get_recolored_icon_png
from engine.components import M3Components
from engine.geometry import generate_scallop_badge_png
from engine.typography import create_textbox, add_styled_paragraph, SCALE_HEADLINE, SCALE_TITLE, SCALE_BODY_SMALL, SCALE_LABEL
from archetypes.base import apply_slide_canvas, render_standard_header, render_standard_footer

def render_team_showcase(slide, theme: M3Theme, data: dict):
    """
    Render Team & Subject Showcase Slide.
    
    Data Schema:
      - category, title, subtitle
      - members: list of 3-4 dicts:
          - name: "Alex Chen"
          - role: "CHIEF ARCHITECT"
          - icon: "psychology"
          - desc: "负责 Google M3 Expressive 底层渲染管线与动态色彩算法工程落地。"
          - tags: ["系统架构", "算法优化", "DrawingML"]
    """
    apply_slide_canvas(slide, theme, style="tinted", add_ambient_bloom=True)

    render_standard_header(
        slide, theme,
        category_text=data.get("category", "ORGANIZATION & TALENT"),
        title_text=data.get("title", "核心专家团队与研发主体"),
        subtitle_text=data.get("subtitle", "跨学科专家团队共同驱动底层算法、渲染内核与设计规范的深度融合"),
        category_icon=data.get("category_icon", "diversity_3")
    )

    members = data.get("members", [
        {
            "name": "Dr. Elena Vance", "role": "CHIEF ARCHITECT", "icon": "psychology",
            "desc": "Google DeepMind 先进编码团队专家，主导 M3 Expressive HCT 感知色彩与 DrawingML 动态编译引擎。",
            "tags": ["HCT 算法", "DrawingML", "管线架构", "动态着色"]
        },
        {
            "name": "Marcus Zhang", "role": "DESIGN SYSTEMS LEAD", "icon": "palette",
            "desc": "10 年前沿 UI/UX 系统专家，负责 28dp Bento 容器美学、有机异形花瓣与排版规范设计。",
            "tags": ["Bento 架构", "M3 Tokens", "无衬线排版", "几何倒角"]
        },
        {
            "name": "Sarah Lin", "role": "ENGINEERING PRINCIPAL", "icon": "account_tree",
            "desc": "专注于高性能跨平台文档编译、矢量图标动态重着色管线与并发安全保护系统构建。",
            "tags": ["Rust 渲染", "矢量引擎", "工业化落地", "跨端编译"]
        }
    ])

    num_members = len(members)
    total_w = 11.733
    start_x = 0.8
    gap = 0.24
    col_w = (total_w - ((num_members - 1) * gap)) / num_members
    card_y = 2.05
    card_h = 4.7

    for i, m in enumerate(members):
        x = start_x + (i * (col_w + gap))
        
        card_res = M3Components.create_card(
            slide, theme,
            left=Inches(x), top=Inches(card_y),
            width=Inches(col_w), height=Inches(card_h),
            variant="elevated", color_role="surface_container_lowest", corner_dp=28, padding=0.28
        )
        inner = card_res.inner_bounds

        # Avatar container with Scallop badge
        scallop_png = f"assets/icons/rendered/team_scallop_{i}.png"
        generate_scallop_badge_png(scallop_png, size=256, petals=12, fill_hex=theme.hex("primary_container"))
        slide.shapes.add_picture(scallop_png, inner.left, inner.top, Inches(0.85), Inches(0.85))

        # Icon inside avatar
        icon_img = get_recolored_icon_png(m.get("icon", "person"), theme.hex("primary"), size=180)
        slide.shapes.add_picture(icon_img, inner.left + Inches(0.2), inner.top + Inches(0.2), Inches(0.45), Inches(0.45))

        # Name
        tb_name = create_textbox(slide, inner.left, inner.top + Inches(1.02), inner.width, Inches(0.45))
        add_styled_paragraph(tb_name.text_frame, m.get("name", "Name"), size_pt=SCALE_HEADLINE, bold=True,
                             color_rgb=theme.rgb("on_surface"))

        # Role Pill
        M3Components.create_chip(
            slide, theme,
            left=inner.left, top=inner.top + Inches(1.52),
            text=m.get("role", "ROLE"), icon_name="verified",
            color_role="tertiary_container", on_color_role="on_tertiary_container", height=0.30
        )

        # Bio / Description
        tb_bio = create_textbox(slide, inner.left, inner.top + Inches(1.95), inner.width, Inches(1.20))
        add_styled_paragraph(tb_bio.text_frame, m.get("desc", ""), size_pt=SCALE_BODY_SMALL,
                             color_rgb=theme.rgb("on_surface_variant"), space_before_pt=2)

        # Skill Tags: 2x2 Grid Layout preventing overflow and eliminating bottom whitespace
        tags = m.get("tags", [])
        chip_gap_emu = Inches(0.12)
        chip_w_emu = (inner.width - chip_gap_emu) / 2.0
        chip_w_val = chip_w_emu / 914400.0
        chip_h = 0.28
        tag_row_y1 = inner.top + Inches(3.28)
        tag_row_y2 = inner.top + Inches(3.64)

        for idx, tag in enumerate(tags[:4]):
            row = idx // 2
            col = idx % 2
            
            if len(tags) == 3 and idx == 2:
                # If only 3 tags, center the 3rd tag on row 2
                t_left = inner.left + (inner.width - chip_w_emu) / 2
                t_top = tag_row_y2
            else:
                t_left = inner.left + col * (chip_w_emu + chip_gap_emu)
                t_top = tag_row_y1 if row == 0 else tag_row_y2

            M3Components.create_chip(
                slide, theme,
                left=t_left, top=t_top,
                text=tag, icon_name=None,
                color_role="surface_container_high", on_color_role="on_surface",
                height=chip_h, width=chip_w_val
            )

    render_standard_footer(slide, theme, left_meta=data.get("footer_left", "Team & Talent Showcase · M3 Expressive"), right_meta=data.get("footer_right", ""))
