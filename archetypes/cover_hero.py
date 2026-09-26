"""
Archetype 01: Hero Keynote Cover Slide.
Designed for strategic keynotes, product launches, and executive decks.
"""

from pptx.util import Inches
from core.tokens import M3Theme
from core.icon_manager import get_recolored_icon_png
from engine.components import M3Components
from engine.geometry import generate_scallop_badge_png
from engine.typography import create_textbox, add_styled_paragraph, SCALE_DISPLAY_LARGE, SCALE_HEADLINE, SCALE_BODY
from archetypes.base import apply_slide_canvas, render_standard_footer

def render_cover_hero(slide, theme: M3Theme, data: dict):
    """
    Render Hero Keynote Cover.
    
    Data Schema:
      - category: str (e.g. "PRODUCT LAUNCH 2026")
      - title: str (e.g. "下一代 AI 智能操作系统与体验设计")
      - subtitle: str (e.g. "基于 Material 3 Expressive 规范的全场景自适应架构")
      - speaker: str (e.g. "Google DeepMind 团队")
      - date: str (e.g. "2026 年秋季发布会")
      - org: str (e.g. "UX & Engineering Architecture")
      - hero_highlight: str (e.g. "EXP-2026")
      - hero_badge_text: str (e.g. "旗舰架构全面演进")
    """
    # 1. Canvas with ambient bloom
    apply_slide_canvas(slide, theme, style="tinted", add_ambient_bloom=True)

    # 2. Category Pill
    category = data.get("category", "KEYNOTE PRESENTATION")
    M3Components.create_chip(
        slide, theme,
        left=Inches(1.0), top=Inches(1.3),
        text=category.upper(),
        icon_name="ai",
        color_role="tertiary_container",
        on_color_role="on_tertiary_container",
        height=0.38
    )

    # 3. Main Title Box
    title_text = data.get("title", "演示文稿核心主标题")
    subtitle_text = data.get("subtitle", "基于 Material 3 Expressive 的全场景演示设计规范")
    
    tb_title = create_textbox(slide, Inches(1.0), Inches(1.85), Inches(6.8), Inches(2.6))
    tf = tb_title.text_frame
    add_styled_paragraph(tf, title_text, size_pt=34, bold=True, color_rgb=theme.rgb("on_surface"))
    add_styled_paragraph(tf, subtitle_text, size_pt=14, bold=False, 
                         color_rgb=theme.rgb("on_surface_variant"), space_before_pt=8)

    # 4. Right Side: Expressive Hero Card with Scallop Badge
    c_right = M3Components.create_card(
        slide, theme,
        left=Inches(8.2), top=Inches(1.3),
        width=Inches(4.1), height=Inches(4.8),
        variant="filled", color_role="primary_container", corner_dp=32
    )

    # Scallop Flower Badge
    scallop_png = "assets/icons/rendered/cover_scallop.png"
    generate_scallop_badge_png(scallop_png, size=320, petals=12, fill_hex=theme.hex("primary"))
    slide.shapes.add_picture(scallop_png, Inches(9.55), Inches(1.7), Inches(1.4), Inches(1.4))

    # White icon inside scallop
    icon_sparkle = get_recolored_icon_png("ai", "#FFFFFF", size=256)
    slide.shapes.add_picture(icon_sparkle, Inches(9.85), Inches(2.0), Inches(0.8), Inches(0.8))

    # Hero Card Text
    tb_hero = create_textbox(slide, Inches(8.5), Inches(3.4), Inches(3.5), Inches(2.3))
    tf_h = tb_hero.text_frame
    hero_code = data.get("hero_highlight", "M3 EXPRESSIVE")
    add_styled_paragraph(tf_h, hero_code, size_pt=18, bold=True, 
                         color_rgb=theme.rgb("primary"), align=None)
    
    hero_badge = data.get("hero_badge_text", "突破低对比度局限 · 激发视觉表现力")
    add_styled_paragraph(tf_h, hero_badge, size_pt=12, bold=False, 
                         color_rgb=theme.rgb("on_primary_container"), space_before_pt=4)
    
    # Mini elevated tag inside hero card
    M3Components.create_chip(
        slide, theme,
        left=Inches(8.5), top=Inches(4.9),
        text="OFFICIAL EDITION", icon_name="verified",
        color_role="surface_container_lowest", on_color_role="primary", height=0.32
    )

    # 5. Bottom Metadata Pills
    meta_y = Inches(5.1)
    speaker = data.get("speaker", "主讲人 · 架构专家")
    date_str = data.get("date", "2026 年秋季")
    org = data.get("org", "Google Technology Group")

    curr_x = Inches(1.0)
    curr_y = meta_y
    for text_val, icon, c_role, oc_role in [
        (speaker, "user", "surface_container_highest", "on_surface"),
        (date_str, "flag", "surface_container_high", "on_surface_variant"),
        (org, "global", "surface_container_high", "on_surface_variant")
    ]:
        zh_count = sum(1 for c in text_val if '\u4e00' <= c <= '\u9fff')
        en_count = len(text_val) - zh_count
        est_w = 0.32 + (en_count * 0.088) + (zh_count * 0.14) + (0.28 if icon else 0)
        if curr_x + Inches(est_w) > Inches(7.8):
            curr_x = Inches(1.0)
            curr_y += Inches(0.44)
        chip = M3Components.create_chip(slide, theme, left=curr_x, top=curr_y, text=text_val,
                                        icon_name=icon, color_role=c_role, on_color_role=oc_role, height=0.34)
        curr_x += chip.width + Inches(0.18)

    # 6. Standard Footer
    render_standard_footer(slide, theme, left_meta="Material 3 Expressive Deck · Cover Hero", right_meta="Page 01")
