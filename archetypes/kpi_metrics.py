"""
Archetype 05: KPI Metrics Dashboard Slide.
High-impact quantitative data wall featuring hero stats, trend indicators, and takeaway banner.
"""

from pptx.util import Inches
from core.tokens import M3Theme
from core.icon_manager import get_recolored_icon_png
from engine.components import M3Components
from engine.typography import create_textbox, add_styled_paragraph, SCALE_DISPLAY_LARGE, SCALE_TITLE, SCALE_BODY, SCALE_BODY_SMALL, MSO_ANCHOR
from archetypes.base import apply_slide_canvas, render_standard_header, render_standard_footer

def render_kpi_metrics(slide, theme: M3Theme, data: dict):
    """
    Render KPI Metrics Dashboard.
    
    Data Schema:
      - category, title, subtitle
      - metrics: list of 3-4 dicts:
          - value: "99.98%"
          - label: "系统可用性 (SLA)"
          - trend: "+0.15% 稳定达标"
          - icon: "security"
          - desc: "高可用架构支撑全天候稳定运行"
          - highlight: bool (if True, uses primary_container)
      - takeaway_text: "综合成果：架构现代化转型实现性能翻倍与研发效能跨越式提升。"
      - takeaway_badge: "KEY TAKEAWAY"
    """
    apply_slide_canvas(slide, theme, style="tinted", add_ambient_bloom=True)

    render_standard_header(
        slide, theme,
        category_text=data.get("category", "KEY PERFORMANCE INDICATORS"),
        title_text=data.get("title", "核心业务指标与成果大屏"),
        subtitle_text=data.get("subtitle", "基于多维量化指标体系的运行效能与业务价值交付总览"),
        category_icon=data.get("category_icon", "analytics")
    )

    metrics = data.get("metrics", [
        {
            "value": "99.98%", "label": "系统高可用性 (SLA)", "trend": "+0.18% 达标",
            "icon": "shield", "desc": "多活容灾与自动化运维保障", "highlight": False
        },
        {
            "value": "12.5 ms", "label": "端到端推理时延", "trend": "-64% 延迟压缩",
            "icon": "speed", "desc": "底层计算图融合与量化加速", "highlight": True
        },
        {
            "value": "4.8x", "label": "全链路研发吞吐", "trend": "+380% 交付提速",
            "icon": "rocket", "desc": "组件化管线带来效能跃迁", "highlight": False
        },
        {
            "value": "86.4%", "label": "用户体验满意度 (CSAT)", "trend": "+24.6% 提升",
            "icon": "star", "desc": "全新 M3 表现力设计语言获高度好评", "highlight": False
        }
    ])

    num_metrics = len(metrics)
    total_w = 11.733
    start_x = 0.8
    gap = 0.22
    col_w = (total_w - ((num_metrics - 1) * gap)) / num_metrics
    cards_y = 2.05
    card_h = 3.2

    for i, m in enumerate(metrics):
        x = start_x + (i * (col_w + gap))
        is_highlight = m.get("highlight", False)

        variant = "filled" if is_highlight else "elevated"
        bg_role = "primary_container" if is_highlight else "surface_container_lowest"
        text_role = "on_primary_container" if is_highlight else "on_surface"
        accent_role = "primary" if is_highlight else "primary"

        card_res = M3Components.create_card(
            slide, theme,
            left=Inches(x), top=Inches(cards_y),
            width=Inches(col_w), height=Inches(card_h),
            variant=variant, color_role=bg_role, corner_dp=24, padding=0.22
        )
        inner = card_res.inner_bounds

        # Tonal Icon Container at top right of card
        icon_bg = "primary" if is_highlight else "surface_container_high"
        icon_fg = "on_primary" if is_highlight else accent_role
        M3Components.create_icon_container(
            slide, theme,
            left=inner.left + inner.width - Inches(0.48),
            top=inner.top,
            size=0.44,
            icon_name=m.get("icon", "insights"),
            color_role=icon_bg,
            icon_color_role=icon_fg,
            shape_type="circle"
        )

        # Big Metric Number
        tb = create_textbox(slide, inner.left, inner.top, inner.width - Inches(0.55), Inches(0.80))
        add_styled_paragraph(tb.text_frame, m.get("value", "0"), size_pt=34, bold=True,
                             color_rgb=theme.rgb(accent_role))

        # Label (13.5pt bold to ensure longer titles stay on a clean single line)
        tb_lbl = create_textbox(slide, inner.left, inner.top + Inches(0.85), inner.width, Inches(0.45))
        add_styled_paragraph(tb_lbl.text_frame, m.get("label", ""), size_pt=13.5, bold=True,
                             color_rgb=theme.rgb(text_role))

        # Trend Chip
        trend_str = m.get("trend", "")
        if trend_str:
            is_pos = "+" in trend_str or "↑" in trend_str or "-" in trend_str
            chip_bg = "tertiary_container" if is_highlight else "primary_container"
            chip_fg = "on_tertiary_container" if is_highlight else "on_primary_container"
            M3Components.create_chip(
                slide, theme,
                left=inner.left, top=inner.top + Inches(1.42),
                text=trend_str, icon_name="growth",
                color_role=chip_bg, on_color_role=chip_fg, height=0.34
            )

        # Bottom Description
        desc_color = "on_primary_container" if is_highlight else "on_surface_variant"
        tb_desc = create_textbox(slide, inner.left, inner.top + Inches(1.88), inner.width, Inches(0.65))
        add_styled_paragraph(tb_desc.text_frame, m.get("desc", ""), size_pt=SCALE_BODY_SMALL,
                             color_rgb=theme.rgb(desc_color), space_before_pt=2)

    # -------------------------------------------------------------
    # Bottom Takeaway Summary Banner
    # -------------------------------------------------------------
    banner_y = cards_y + card_h + 0.24
    banner_h = 1.15
    banner_res = M3Components.create_card(
        slide, theme,
        left=Inches(start_x), top=Inches(banner_y),
        width=Inches(total_w), height=Inches(banner_h),
        variant="filled", color_role="surface_container", corner_dp=24, padding=0.2
    )
    b_inner = banner_res.inner_bounds

    # Badge on left of banner
    takeaway_badge = data.get("takeaway_badge", "KEY TAKEAWAY")
    badge_h = 0.36
    badge_top = b_inner.top + (b_inner.height - Inches(badge_h)) / 2
    badge_chip = M3Components.create_chip(
        slide, theme,
        left=b_inner.left, top=badge_top,
        text=takeaway_badge, icon_name="lightbulb",
        color_role="tertiary_container", on_color_role="on_tertiary_container", height=badge_h
    )

    # Dynamic text positioning starting after badge width + spacing
    badge_w_in = badge_chip.width / 914400.0
    gap = 0.22
    text_left = b_inner.left + Inches(badge_w_in + gap)
    text_width = b_inner.width - Inches(badge_w_in + gap)

    takeaway_text = data.get("takeaway_text", "综合成效：通过架构现代化改造与全链路组件化升级，系统关键 SLA 指标全面达成，用户体验与业务吞吐实现质的跃迁。")
    tb_b = create_textbox(slide, text_left, b_inner.top, text_width, b_inner.height, vertical_anchor=MSO_ANCHOR.MIDDLE)
    add_styled_paragraph(tb_b.text_frame, takeaway_text, size_pt=12, bold=False,
                         color_rgb=theme.rgb("on_surface"))

    render_standard_footer(slide, theme, left_meta=data.get("footer_left", "KPI Metrics Dashboard · M3 Expressive"), right_meta=data.get("footer_right", ""))
