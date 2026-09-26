"""
Archetype 03: Process Roadmap / Timeline Slide.
Sequential milestone track with phase pills, status badges, and deliverable cards.
"""

from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from core.tokens import M3Theme
from core.icon_manager import get_recolored_icon_png
from engine.components import M3Components
from engine.typography import create_textbox, add_styled_paragraph, SCALE_TITLE, SCALE_BODY_SMALL, SCALE_LABEL
from archetypes.base import apply_slide_canvas, render_standard_header, render_standard_footer

def render_process_roadmap(slide, theme: M3Theme, data: dict):
    """
    Render Process Roadmap / Timeline.
    
    Data Schema:
      - category, title, subtitle
      - phases: list of 3-4 dicts:
          - phase_num: "01"
          - name: "需求调研与算法验证"
          - time: "Q1 2026"
          - status: "COMPLETED" | "ACTIVE" | "UPCOMING"
          - icon: "check" | "bolt" | "rocket"
          - deliverables: ["交付项 1", "交付项 2", "交付项 3"]
          - kpi_chip: "完成率 100%"
    """
    apply_slide_canvas(slide, theme, style="tinted", add_ambient_bloom=True)

    render_standard_header(
        slide, theme,
        category_text=data.get("category", "DEVELOPMENT ROADMAP"),
        title_text=data.get("title", "演进路线图与阶段里程碑"),
        subtitle_text=data.get("subtitle", "四阶段推进计划：从基础设施搭建到全场景工业化落地"),
        category_icon=data.get("category_icon", "flag")
    )

    phases = data.get("phases", [
        {
            "phase_num": "01", "name": "官方资产集成", "time": "阶段一 · 已就绪", 
            "status": "COMPLETED", "icon": "check",
            "deliverables": ["Google MCU 算法接入", "33 款矢量图标本地缓存", "动态 HCT 双模式色阶"],
            "kpi_chip": "基础设施 100%"
        },
        {
            "phase_num": "02", "name": "表现力组件引擎", "time": "阶段二 · 进行中", 
            "status": "ACTIVE", "icon": "bolt",
            "deliverables": ["28dp 核心大圆角规范", "9999px 全胶囊药丸算子", "中西文抗降级排版架构"],
            "kpi_chip": "体验提升 4.8x"
        },
        {
            "phase_num": "03", "name": "8大核心版式库", "time": "阶段三 · 规划中", 
            "status": "UPCOMING", "icon": "layers",
            "deliverables": ["战略封面与便当盒看板", "时间轴与双方案对比", "指标墙与特性矩阵网格"],
            "kpi_chip": "8 组原型覆盖"
        },
        {
            "phase_num": "04", "name": "Skill 封装与验收", "time": "阶段四 · 目标", 
            "status": "UPCOMING", "icon": "rocket",
            "deliverables": ["自然语言大纲解析器", "全局 SKILL.md 注册挂载", "5 页商业级实机验收"],
            "kpi_chip": "端到端交付"
        }
    ])

    num_phases = len(phases)
    total_w = 11.733
    start_x = 0.8
    gap = 0.22
    col_w = (total_w - ((num_phases - 1) * gap)) / num_phases
    card_y = 2.05
    card_h = 4.7

    for i, p in enumerate(phases):
        x = start_x + (i * (col_w + gap))
        status = p.get("status", "UPCOMING")

        # Determine card variant and colors based on status
        if status == "ACTIVE":
            variant = "filled"
            bg_role = "tertiary_container"
            on_role = "on_tertiary_container"
            accent_role = "tertiary"
        elif status == "COMPLETED":
            variant = "filled"
            bg_role = "primary_container"
            on_role = "on_primary_container"
            accent_role = "primary"
        else: # UPCOMING
            variant = "elevated"
            bg_role = "surface_container_lowest"
            on_role = "on_surface"
            accent_role = "primary"

        # Main Phase Card
        card_res = M3Components.create_card(
            slide, theme,
            left=Inches(x), top=Inches(card_y),
            width=Inches(col_w), height=Inches(card_h),
            variant=variant, color_role=bg_role, corner_dp=24, padding=0.22
        )

        inner = card_res.inner_bounds

        # Phase Status Chip at Top
        phase_num_text = f"PHASE {p.get('phase_num', f'0{i+1}')}"
        if status == "ACTIVE":
            chip_bg = "tertiary"
            chip_fg = "on_tertiary"
        elif status == "COMPLETED":
            chip_bg = "primary"
            chip_fg = "on_primary"
        else:
            chip_bg = "surface_container_high"
            chip_fg = "on_surface_variant"

        M3Components.create_chip(
            slide, theme,
            left=inner.left, top=inner.top,
            text=phase_num_text, icon_name=p.get("icon", "flag"),
            color_role=chip_bg, on_color_role=chip_fg, height=0.3
        )

        # Phase Name and Time
        tb_title = create_textbox(slide, inner.left, inner.top + Inches(0.48), inner.width, Inches(0.85))
        add_styled_paragraph(tb_title.text_frame, p.get("name", "阶段名称"), size_pt=SCALE_TITLE, bold=True,
                             color_rgb=theme.rgb(on_role))
        add_styled_paragraph(tb_title.text_frame, p.get("time", "时间节点"), size_pt=10, bold=False,
                             color_rgb=theme.rgb(accent_role), space_before_pt=2)

        # Deliverables header (14pt bold per user specification)
        hdr_top = inner.top + Inches(1.48)
        tb_deliv_hdr = create_textbox(slide, inner.left, hdr_top, inner.width, Inches(0.32))
        add_styled_paragraph(tb_deliv_hdr.text_frame, "核心成果清单:", size_pt=14.0, bold=True,
                             color_rgb=theme.rgb(accent_role))

        # Deliverables checklist with distinct micro-badges (unwrapped)
        card_micro_icons = [
            ["tune", "dns", "palette"],
            ["bolt", "grid_view", "spellcheck"],
            ["layers", "auto_awesome", "palette"],
            ["rocket", "terminal", "verified"]
        ]
        phase_icons = card_micro_icons[i % len(card_micro_icons)]
        
        badge_bg = "primary_container" if status != "ACTIVE" else "tertiary_container"
        badge_fg = "primary" if status != "ACTIVE" else "tertiary"

        deliv_items = []
        for idx, item_text in enumerate(p.get("deliverables", [])):
            icon_for_item = phase_icons[idx % len(phase_icons)]
            deliv_items.append({
                "text": item_text,
                "icon": icon_for_item,
                "bg_role": badge_bg,
                "fg_role": badge_fg
            })

        M3Components.create_checklist(
            slide, theme,
            left=inner.left, top=hdr_top + Inches(0.44),
            width=inner.width,
            items=deliv_items,
            font_size_pt=10.5,
            text_color_role=on_role,
            item_spacing_in=0.46,
            badge_size=0.19
        )

        # Bottom KPI micro capsule: Each card has a distinct icon and tailored tonal style
        kpi_text = p.get("kpi_chip", "阶段成果")
        default_kpi_icons = ["tune", "bolt", "layers", "rocket"]
        kpi_icon = p.get("kpi_icon", default_kpi_icons[i % len(default_kpi_icons)])
        
        if status == "ACTIVE":
            kpi_bg = "surface_container_lowest"
            kpi_fg = "tertiary"
        elif status == "COMPLETED":
            kpi_bg = "surface_container_lowest"
            kpi_fg = "primary"
        else:
            kpi_bg = "primary_container"
            kpi_fg = "primary"

        M3Components.create_chip(
            slide, theme,
            left=inner.left, top=inner.top + Inches(3.82),
            text=kpi_text, icon_name=kpi_icon,
            color_role=kpi_bg,
            on_color_role=kpi_fg, height=0.32
        )

    render_standard_footer(slide, theme, left_meta=data.get("footer_left", "Process Roadmap · M3 Expressive"), right_meta=data.get("footer_right", ""))
