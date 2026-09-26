"""
Archetype 02: Bento Grid Dashboard Slide.
Modular multi-compartment container layout for system overviews and key pillars.
"""

from pptx.util import Inches
from pptx.enum.text import MSO_ANCHOR
from core.tokens import M3Theme
from core.icon_manager import get_recolored_icon_png
from engine.components import M3Components, M3BentoGrid
from engine.geometry import generate_scallop_badge_png
from engine.typography import create_textbox, add_styled_paragraph, SCALE_HEADLINE, SCALE_TITLE, SCALE_BODY, SCALE_BODY_SMALL
from archetypes.base import apply_slide_canvas, render_standard_header, render_standard_footer

def render_bento_grid(slide, theme: M3Theme, data: dict):
    """
    Render Bento Grid Dashboard.
    """
    # 1. Canvas
    apply_slide_canvas(slide, theme, style="tinted", add_ambient_bloom=True)

    # 2. Header
    render_standard_header(
        slide, theme,
        category_text=data.get("category", "SYSTEM OVERVIEW"),
        title_text=data.get("title", "Bento 便当盒全景看板"),
        subtitle_text=data.get("subtitle", "基于模块化分舱架构与 28dp 大圆角设计的全景概览"),
        category_icon=data.get("category_icon", "bento")
    )

    # 3. Grid Partitioning
    b_hero, b_top, b_bot_left, b_bot_right = M3BentoGrid.layout_1_hero_3_sub(
        x=0.8, y=2.05, w=11.733, h=4.7, gap=0.24
    )

    # 4. CARD 1: Hero Card (Left)
    hero_data = data.get("hero_card", {})
    hero_res = M3Components.create_card(
        slide, theme,
        left=Inches(b_hero.left), top=Inches(b_hero.top),
        width=Inches(b_hero.width), height=Inches(b_hero.height),
        variant="filled", color_role="primary_container", corner_dp=28
    )
    
    icon_name = hero_data.get("icon", "ai")
    icon_ai = get_recolored_icon_png(icon_name, theme.hex("primary"), size=256)
    slide.shapes.add_picture(icon_ai, hero_res.inner_bounds.left, hero_res.inner_bounds.top, Inches(0.55), Inches(0.55))

    M3Components.create_chip(
        slide, theme,
        left=hero_res.inner_bounds.left + Inches(0.72), top=hero_res.inner_bounds.top + Inches(0.08),
        text=hero_data.get("chip", "CORE PILLAR"), icon_name="verified",
        color_role="primary", on_color_role="on_primary", height=0.32
    )

    tb_hero = create_textbox(slide, hero_res.inner_bounds.left, hero_res.inner_bounds.top + Inches(0.8), 
                             hero_res.inner_bounds.width, Inches(1.8))
    add_styled_paragraph(tb_hero.text_frame, hero_data.get("title", "核心支柱标题"), 
                         size_pt=SCALE_HEADLINE, bold=True, color_rgb=theme.rgb("on_primary_container"))
    add_styled_paragraph(tb_hero.text_frame, hero_data.get("subtitle", "Subheadline Overview"), 
                         size_pt=11, bold=True, color_rgb=theme.rgb("primary"), space_before_pt=2, space_after_pt=6)
    add_styled_paragraph(tb_hero.text_frame, hero_data.get("desc", "详细核心阐述段落，用于说明系统设计理念与关键技术演进路径。"), 
                         size_pt=SCALE_BODY, bold=False, color_rgb=theme.rgb("on_primary_container"))

    # Hero Nested Metric Box
    metric_val = hero_data.get("metric_val", "84.6%")
    metric_lbl = hero_data.get("metric_label", "关键指标提升率")
    trend = hero_data.get("trend", "+3.2x")
    M3Components.create_hero_metric(
        slide, theme,
        left=hero_res.inner_bounds.left, top=hero_res.inner_bounds.top + Inches(2.7),
        width=hero_res.inner_bounds.width, height=Inches(1.4),
        metric_value=metric_val, metric_label=metric_lbl,
        trend_str=trend, color_role="primary", bg_role="surface_container_lowest"
    )

    # 5. CARD 2: Top Right Card
    top_data = data.get("card_top", {})
    c2_res = M3Components.create_card(
        slide, theme,
        left=Inches(b_top.left), top=Inches(b_top.top),
        width=Inches(b_top.width), height=Inches(b_top.height),
        variant="outlined", color_role="surface_container", corner_dp=28
    )
    M3Components.create_icon_container(
        slide, theme,
        left=c2_res.inner_bounds.left, top=c2_res.inner_bounds.top,
        size=0.48, icon_name=top_data.get("icon", "grid_view"),
        color_role="surface_container_highest", icon_color_role="primary",
        shape_type="circle"
    )

    tb_c2 = create_textbox(slide, c2_res.inner_bounds.left + Inches(0.65), c2_res.inner_bounds.top, 
                           c2_res.inner_bounds.width - Inches(0.65), Inches(0.85))
    add_styled_paragraph(tb_c2.text_frame, top_data.get("title", "模块化组件分舱"), 
                         size_pt=SCALE_TITLE, bold=True, color_rgb=theme.rgb("on_surface"))
    add_styled_paragraph(tb_c2.text_frame, top_data.get("desc", "卡片内外边距严格依循 M3 规范，实现各功能区视觉和谐。"), 
                         size_pt=SCALE_BODY_SMALL, bold=False, color_rgb=theme.rgb("on_surface_variant"), space_before_pt=3)

    chips = top_data.get("chips", ["28dp 容器", "全胶囊药丸", "色调浸润"])
    chip_x = c2_res.inner_bounds.left + Inches(0.65)
    chip_y = c2_res.inner_bounds.top + Inches(1.05)
    for chip_text in chips:
        pill = M3Components.create_chip(
            slide, theme, left=chip_x, top=chip_y,
            text=chip_text, icon_name="bolt", 
            color_role="primary_container", on_color_role="on_primary_container", 
            height=0.30
        )
        chip_x += pill.width + Inches(0.14)

    # 6. CARD 3: Bottom Left Card (Tertiary Accent)
    bot_l_data = data.get("card_bot_left", {})
    c3_res = M3Components.create_card(
        slide, theme,
        left=Inches(b_bot_left.left), top=Inches(b_bot_left.top),
        width=Inches(b_bot_left.width), height=Inches(b_bot_left.height),
        variant="filled", color_role="tertiary_container", corner_dp=24
    )
    
    # Unified 0.44" icon container
    M3Components.create_icon_container(
        slide, theme,
        left=c3_res.inner_bounds.left, top=c3_res.inner_bounds.top,
        size=0.44, icon_name=bot_l_data.get("icon", "bolt"),
        color_role="tertiary", icon_color_role="on_tertiary",
        shape_type="circle"
    )

    # Title beside icon, vertically centered with the icon
    tb_c3_title = create_textbox(
        slide, 
        c3_res.inner_bounds.left + Inches(0.55), 
        c3_res.inner_bounds.top, 
        c3_res.inner_bounds.width - Inches(0.55), 
        Inches(0.44),
        vertical_anchor=MSO_ANCHOR.MIDDLE
    )
    add_styled_paragraph(tb_c3_title.text_frame, bot_l_data.get("title", "表现力强调"), 
                         size_pt=15, bold=True, color_rgb=theme.rgb("on_tertiary_container"))

    # Symmetrical content below: stat metric + description
    tb_c3_stat = create_textbox(
        slide, 
        c3_res.inner_bounds.left, 
        c3_res.inner_bounds.top + Inches(0.58), 
        c3_res.inner_bounds.width, 
        Inches(0.52)
    )
    add_styled_paragraph(tb_c3_stat.text_frame, bot_l_data.get("stat", "+45%"), 
                         size_pt=28, bold=True, color_rgb=theme.rgb("tertiary"))

    tb_c3_desc = create_textbox(
        slide, 
        c3_res.inner_bounds.left, 
        c3_res.inner_bounds.top + Inches(1.15), 
        c3_res.inner_bounds.width, 
        Inches(0.70)
    )
    add_styled_paragraph(tb_c3_desc.text_frame, bot_l_data.get("desc", "通过高光跃迁锁定视线。"), 
                         size_pt=12.0, bold=False, color_rgb=theme.rgb("on_tertiary_container"), space_before_pt=2)

    # 7. CARD 4: Bottom Right Card (Elevated White)
    bot_r_data = data.get("card_bot_right", {})
    c4_res = M3Components.create_card(
        slide, theme,
        left=Inches(b_bot_right.left), top=Inches(b_bot_right.top),
        width=Inches(b_bot_right.width), height=Inches(b_bot_right.height),
        variant="elevated", color_role="surface_container_lowest", corner_dp=24
    )
    
    # Unified 0.44" icon container
    M3Components.create_icon_container(
        slide, theme,
        left=c4_res.inner_bounds.left, top=c4_res.inner_bounds.top,
        size=0.44, icon_name=bot_r_data.get("icon", "verified"),
        color_role="primary_container", icon_color_role="primary",
        shape_type="circle"
    )

    # Title beside icon, vertically centered with the icon
    tb_c4_title = create_textbox(
        slide, 
        c4_res.inner_bounds.left + Inches(0.55), 
        c4_res.inner_bounds.top, 
        c4_res.inner_bounds.width - Inches(0.55), 
        Inches(0.44),
        vertical_anchor=MSO_ANCHOR.MIDDLE
    )
    add_styled_paragraph(tb_c4_title.text_frame, bot_r_data.get("title", "技术指标"), 
                         size_pt=13.5, bold=True, color_rgb=theme.rgb("on_surface"))
    
    # Checklist with increased font size and line spacing to balance whitespace
    bullets = bot_r_data.get("bullets", ["纯原生矢量对象", "抗字体降级体系", "100% 自由编辑修改"])
    M3Components.create_checklist(
        slide, theme,
        left=c4_res.inner_bounds.left, 
        top=c4_res.inner_bounds.top + Inches(0.58),
        width=c4_res.inner_bounds.width,
        items=bullets,
        icon_name="verified",
        icon_bg_role="primary_container",
        icon_fg_role="primary",
        font_size_pt=11.5,
        text_color_role="on_surface",
        item_spacing_in=0.38,
        badge_size=0.20
    )

    # 8. Footer
    render_standard_footer(slide, theme, left_meta=data.get("footer_left", "Bento Grid Dashboard · M3 Expressive"), right_meta=data.get("footer_right", ""))
