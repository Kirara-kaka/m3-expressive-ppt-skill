"""
Common slide layout helpers and header/footer components for M3 Archetypes.
"""

from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from core.tokens import M3Theme
from engine.components import M3Components, M3Canvas
from engine.typography import (
    create_textbox, 
    add_styled_paragraph, 
    SCALE_DISPLAY_MEDIUM, 
    SCALE_BODY
)

def apply_slide_canvas(slide, theme: M3Theme, style: str = "tinted", add_ambient_bloom: bool = True):
    """Apply unified background canvas to the slide."""
    return M3Canvas.apply_background(slide, theme, style=style, add_ambient_bloom=add_ambient_bloom)

def render_standard_header(slide, 
                           theme: M3Theme, 
                           category_text: str, 
                           title_text: str, 
                           subtitle_text: str = "",
                           category_icon: str = "ai",
                           top_in: float = 0.52):
    """
    Render consistent, elegant M3 Expressive slide header:
    Category pill chip + Display title + Secondary subtitle.
    """
    # Category Pill
    if category_text:
        M3Components.create_chip(
            slide, theme,
            left=Inches(0.8), top=Inches(top_in),
            text=category_text.upper(),
            icon_name=category_icon,
            color_role="tertiary_container",
            on_color_role="on_tertiary_container",
            height=0.36
        )

    # Title & Subtitle Box
    title_top = top_in + (0.44 if category_text else 0)
    tb = create_textbox(slide, Inches(0.8), Inches(title_top), Inches(11.733), Inches(0.85))
    tf = tb.text_frame
    
    add_styled_paragraph(tf, title_text, size_pt=SCALE_DISPLAY_MEDIUM, bold=True, 
                         color_rgb=theme.rgb("on_surface"))
    
    if subtitle_text:
        add_styled_paragraph(tf, subtitle_text, size_pt=11.5, bold=False, 
                             color_rgb=theme.rgb("on_surface_variant"), space_before_pt=3)

def render_standard_footer(slide, 
                           theme: M3Theme, 
                           left_meta: str = "Material 3 Expressive Specification · Automated Deck",
                           right_meta: str = ""):
    """Render subtle metadata footer at bottom of slide."""
    tb = create_textbox(slide, Inches(0.8), Inches(6.95), Inches(11.733), Inches(0.35))
    tf = tb.text_frame
    p = add_styled_paragraph(tf, left_meta, size_pt=8.5, color_rgb=theme.rgb("outline"))
    
    if right_meta:
        tb_r = create_textbox(slide, Inches(6.0), Inches(6.95), Inches(6.533), Inches(0.35))
        add_styled_paragraph(tb_r.text_frame, right_meta, size_pt=8.5, 
                             color_rgb=theme.rgb("outline"), align=PP_ALIGN.RIGHT)
