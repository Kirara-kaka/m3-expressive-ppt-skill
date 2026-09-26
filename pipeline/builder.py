"""
Deck Builder Engine for M3 Expressive presentation generation.
Compiles DeckPlan into native, editable DrawingML PPTX presentations.
"""

import sys
import os
import json
import argparse
from typing import Dict, Any, Union

from pptx import Presentation
from pptx.util import Inches

from core.tokens import get_preset_theme, generate_theme, M3Theme
from archetypes import (
    render_cover_hero,
    render_bento_grid,
    render_process_roadmap,
    render_dual_contrast,
    render_kpi_metrics,
    render_feature_list,
    render_quote_takeaway,
    render_team_showcase,
)
from pipeline.schema import DeckPlan
from pipeline.planner import auto_plan_from_outline

ARCHETYPE_RENDERERS = {
    "cover_hero": render_cover_hero,
    "bento_grid": render_bento_grid,
    "process_roadmap": render_process_roadmap,
    "dual_contrast": render_dual_contrast,
    "kpi_metrics": render_kpi_metrics,
    "feature_list": render_feature_list,
    "quote_takeaway": render_quote_takeaway,
    "team_showcase": render_team_showcase,
}

def safe_save_pptx(prs: Presentation, target_path: str) -> str:
    """Saves presentation, auto-generating _v1, _v2 if file is locked by PowerPoint."""
    os.makedirs(os.path.dirname(os.path.abspath(target_path)), exist_ok=True)
    try:
        prs.save(target_path)
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        for i in range(1, 25):
            cand = f"{base}_v{i}{ext}"
            try:
                prs.save(cand)
                print(f"Notice: {target_path} is locked by PowerPoint. Saved to {cand} instead.")
                return cand
            except PermissionError:
                continue
        raise PermissionError(f"Could not save {target_path} in any version.")

class M3DeckBuilder:
    """Assembles and generates a full presentation deck from a DeckPlan."""

    def __init__(self, plan: DeckPlan):
        self.plan = plan
        self.prs = Presentation()
        self._configure_dimensions()
        self.theme = self._resolve_theme()

    def _configure_dimensions(self):
        # 16:9 standard widescreen presentation (13.333 x 7.5 inches)
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def _resolve_theme(self) -> M3Theme:
        cfg = self.plan.config
        if cfg.seed_color:
            return generate_theme(cfg.seed_color, is_dark=cfg.is_dark)
        return get_preset_theme(cfg.preset_theme, is_dark=cfg.is_dark)

    def build(self, output_path: str = "output.pptx") -> str:
        blank_layout = self.prs.slide_layouts[6]
        
        print(f"Generating Presentation: '{self.plan.config.title}' ({len(self.plan.slides)} slides)...")
        print(f"Theme: {self.plan.config.preset_theme} (Dark: {self.plan.config.is_dark})")

        for idx, slide_plan in enumerate(self.plan.slides, start=1):
            slide = self.prs.slides.add_slide(blank_layout)
            renderer = ARCHETYPE_RENDERERS.get(slide_plan.archetype)
            if not renderer:
                print(f"Warning: Archetype '{slide_plan.archetype}' not found. Falling back to bento_grid.")
                renderer = render_bento_grid

            renderer(slide, self.theme, slide_plan.content)
            print(f"  [OK] Page {idx:02d}/{len(self.plan.slides):02d} rendered ({slide_plan.archetype})")

        saved_path = safe_save_pptx(self.prs, output_path)
        print(f"[PASS] Successfully saved presentation to: {saved_path}")
        return saved_path

def build_deck_from_dict(data: Dict[str, Any], output_path: str = "output.pptx") -> str:
    plan = auto_plan_from_outline(data)
    builder = M3DeckBuilder(plan)
    return builder.build(output_path)

def build_deck_from_json(json_input: Union[str, Dict[str, Any]], output_path: str = "output.pptx") -> str:
    if isinstance(json_input, str):
        if os.path.exists(json_input):
            with open(json_input, "r", encoding="utf-8") as f:
                data = json.load(f)
        else:
            data = json.loads(json_input)
    else:
        data = json_input

    return build_deck_from_dict(data, output_path)

def main():
    parser = argparse.ArgumentParser(description="M3 Expressive Deck Builder CLI")
    parser.add_argument("input", help="Path to JSON outline file or JSON string")
    parser.add_argument("-o", "--output", default="output.pptx", help="Output PPTX path")
    args = parser.parse_args()

    out_file = build_deck_from_json(args.input, args.output)
    print(f"Done: {out_file}")

if __name__ == "__main__":
    main()
